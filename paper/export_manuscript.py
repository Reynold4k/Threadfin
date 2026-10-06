#!/usr/bin/env python3
"""Render the reviewed source with template styles, numbered citations and RIS.

This exports existing results only. It never runs a biological analysis.
Final export requires every planned figure and the completed native comparison.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
PAPER = ROOT / 'paper'
FIGURES = PAPER / 'figure_plan'
TEMPLATE = PROJECT.parent / '文献调研/演讲文稿/Progress_Review_Year2_ChenZhu_v5.docx'
BENCH = ROOT / 'case_studies/results/native_benchmark'


def native_completed():
    return json.loads((BENCH/'pipeline_status.json').read_text()).get('status') == 'completed'


def citation(record):
    authors = record['authors']
    def author(a):
        given=a.get('given','')
        initials=''.join(re.findall(r'(?:^|[\s-])([A-Za-z])',given))
        return a['family'] + (' ' + initials if initials else '')
    names=', '.join(author(a) for a in authors[:6])
    if len(authors)>6:names+=', et al'
    journal=record['journal_abbrev']
    if record['status']=='preprint':journal+=' [Preprint]'
    volume=record.get('volume') or ''
    if volume and record.get('issue'):volume+=f"({record['issue']})"
    pagination=(';' + volume if volume else '')
    if record.get('pages'):pagination+=(':' if volume else ';') + record['pages']
    return f"{names}. {record['title'].rstrip('.')}. {journal}. {record['year']}{pagination}. doi:{record['doi']}."


def existing_rbd_result(allow_pending):
    summary=BENCH/'readout_summary.csv'
    state=json.loads((BENCH/'pipeline_status.json').read_text())
    if state.get('status')!='completed' or not summary.exists():
        if not allow_pending:raise RuntimeError('RBD native comparison is incomplete; final manuscript export is deferred.')
        return ('The complete RBD native comparison is still pending. No RBD performance '\
                'is inferred from NP results or partial representations. Figure 6 and '\
                'Supplementary Figure 7 await all native outputs and visual review.')
    rows=list(csv.DictReader(summary.open()))
    selected=[r for r in rows if r['dataset']=='mouse_rbd' and r['target']=='division_gate:mCherry-low' and int(r['min_cells'])==2]
    by={r['method']:r for r in selected}
    methods=['RNA_centroid','RNA_context_centroid','Threadfin_mean','Threadfin_kernel','Benisse','BiGCN','clone2vec','Training_mean']
    if set(by)!=set(methods):raise RuntimeError('RBD saved summary lacks the complete comparator set.')
    def mae(method):return f"{float(by[method]['median_mae']):.3f}"
    mice={int(r['scored_mice']) for r in selected};families={int(r['n_scored_families']) for r in selected}
    if len(mice)!=1 or len(families)!=1:raise RuntimeError('Saved RBD comparators use inconsistent fold coverage.')
    return (f"The completed RBD comparison retains {families.pop():,} reporter-eligible families "
            f"across {mice.pop()} label-held-out mice. Median mouse-wise absolute error is "
            f"{mae('Threadfin_mean')} for Threadfin mean profiles and {mae('Threadfin_kernel')} "
            f"for kernel profiles; raw and donor-centred RNA means yield {mae('RNA_centroid')} "
            f"and {mae('RNA_context_centroid')}. Benisse, BiGCN and clone2vec yield "
            f"{mae('Benisse')}, {mae('BiGCN')} and {mae('clone2vec')}, respectively, compared "
            f"with {mae('Training_mean')} for the training-mean control. These values describe "
            "the prespecified captured reporter-fraction readout, not an overall tool ranking.")


def render_source(allow_pending):
    refs=json.loads((PAPER/'references.json').read_text())
    source=(PAPER/'MANUSCRIPT_source.md').read_text()
    source=source.replace('{RBD_BENCHMARK_RESULTS}',existing_rbd_result(allow_pending))
    if not native_completed():
        source=source.replace('Draft 3, 6 October 2026.',
            'Review draft, 6 October 2026. RBD native benchmark and Figure 6/S7 remain pending.')
    order=[]
    def replace(match):
        numbers=[]
        for key in match[1].split(','):
            if key not in refs:raise KeyError(f'Unregistered citation: {key}')
            if key not in order:order.append(key)
            numbers.append(str(order.index(key)+1))
        return '('+', '.join(numbers)+')'
    source=re.sub(r'\{cite:([^}]+)\}',replace,source)
    bibliography='\n\n'.join(f'{i}. {citation(refs[key])}' for i,key in enumerate(order,1))
    source=source.replace('{REFERENCES}',bibliography)
    if re.search(r'\{(?:cite:|RBD_|REFERENCES)',source):raise ValueError('Unresolved manuscript placeholder')
    return source.rstrip(),refs,order


def hyperlink(paragraph,text,url):
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');prop=OxmlElement('w:rPr');style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');prop.append(style)
    run.append(prop);el=OxmlElement('w:t');el.text=text;run.append(el);link.append(run);paragraph._p.append(link)


def inline(paragraph,text,italic=False):
    pattern=r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))'
    for token in re.split(pattern,text):
        if not token:continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',token)
        if m:
            url=m[2]
            if not re.match(r'https?://',url):url='https://github.com/Reynold4k/Threadfin/blob/manuscript/gc-biology/paper/'+url
            hyperlink(paragraph,m[1],url);continue
        run=paragraph.add_run(token.strip('*`') if token[0] in '*`' else token)
        run.italic=italic or (token.startswith('*') and not token.startswith('**'))
        if token.startswith('**'):run.bold=True
        if token.startswith('`'):run.font.name='Courier New';run.font.size=Pt(10)


def fresh_template(template):
    original=Document(template);doc=Document()
    doc.part._styles_part._element=deepcopy(original.styles.element)
    # Preserve theme/fonts but not unrelated report text, images or EndNote fields.
    for reltype in [RT.THEME,RT.FONT_TABLE]:
        old=next((r.target_part for r in original.part.rels.values() if r.reltype==reltype),None)
        new=next((r.target_part for r in doc.part.rels.values() if r.reltype==reltype),None)
        if old is not None and new is not None:new._blob=old.blob
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Threadfin  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(9)
    doc.core_properties.title='Threadfin links B-cell receptor families to germinal-centre cell-state distributions'
    doc.core_properties.author='Chen Zhu'
    doc.core_properties.subject='GC-focused manuscript; existing public-data results'
    return doc,original


def add_figure(doc,name,legend,missing_figures):
    doc.add_page_break()
    image=FIGURES/(name+'.png')
    if name in missing_figures:
        doc.add_paragraph(name.replace('_',' ')+' — awaiting complete native benchmark',style='Heading 2')
    else:
        with Image.open(image) as bitmap:width,height=bitmap.size
        scale=min(Cm(17)/width,Cm(24.2)/height)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        shape=p.add_run().add_picture(str(image),width=int(width*scale),height=int(height*scale))
        shape._inline.docPr.set('descr',legend['title'])
        # Keep the source image intact and place a readable full legend on the next page.
        doc.add_page_break()
    paragraph=doc.add_paragraph(style='FigureLegend')
    paragraph.add_run(legend['title']+'. ').bold=True
    inline(paragraph,legend['body'])


def legends():
    text=(PAPER/'FIGURE_LEGENDS.md').read_text();result={}
    for title,body in re.findall(r'^## ([^\n]+)\n\n(.*?)(?=\n## |\Z)',text,flags=re.M|re.S):
        match=re.match(r'(Supplementary )?Figure (\d+) \| (.+)',title)
        if not match:continue
        kind='Supplementary' if match[1] else 'Figure';key=kind+'_'+match[2]
        result[key]={'title':kind+' Figure '+match[2]+' '+match[3] if kind=='Supplementary' else 'Figure '+match[2]+' '+match[3],
                     'body':' '.join(body.split())}
    return result


def ris(refs,order):
    records=[]
    for key in order:
        r=refs[key];lines=['TY  - JOUR' if r['status']=='published' else 'TY  - UNPB']
        lines += ['AU  - '+a['family']+', '+a.get('given','') for a in r['authors']]
        lines += ['TI  - '+r['title'],'JO  - '+r['journal'],'JA  - '+r['journal_abbrev'],'PY  - '+str(r['year'])]
        for tag,value in [('VL',r.get('volume')),('IS',r.get('issue')),('SP',r.get('pages')),('DO',r['doi']),('UR',r['url'])]:
            if value:lines.append(tag+'  - '+value)
        if r['status']=='preprint':lines.append('N1  - Preprint; not peer reviewed')
        records.append('\n'.join(lines+['ER  -','']))
    return '\n'.join(records)


def validate_docx(doc,template,output,order,missing_figures,template_path):
    from zipfile import ZipFile
    imported=['a1','1','21','31','FigureLegend','EndNoteBibliography']
    for sid in imported:
        left=doc.styles.element.xpath(f'./w:style[@w:styleId="{sid}"]')[0]
        right=template.styles.element.xpath(f'./w:style[@w:styleId="{sid}"]')[0]
        assert ET.tostring(ET.fromstring(left.xml))==ET.tostring(ET.fromstring(right.xml)),sid
    with ZipFile(output) as z:
        xml=z.read('word/document.xml').decode()
        assert 'ADDIN EN.CITE' not in xml  # no invalid borrowed EndNote records
        assert 'Chen Zhu' in xml and 'Peter Doherty Institute' in xml
        assert '{RBD_BENCHMARK_RESULTS}' not in xml and '{cite:' not in xml
        images=[n for n in z.namelist() if n.startswith('word/media/')]
        assert len(images)==14-len(missing_figures)
    return {'template':str(template_path),'template_sha256':hashlib.sha256(template_path.read_bytes()).hexdigest(),
            'output':str(output),'author':'Chen Zhu','styles_preserved':imported,'citation_count':len(order),
            'inline_citations':'parenthetical numbers in first-citation order',
            'reference_management':'static Word citations plus EndNote-importable RIS; no fabricated EndNote fields',
            'figure_images':len(images),'pending_native_benchmark':not native_completed(),
            'missing_figures':missing_figures,
            'validation':'OOXML/styles/media/reference checks; Word pagination requires visual review'}


def export(output,allow_pending=False,template_path=TEMPLATE):
    source,refs,order=render_source(allow_pending)
    expected=[f'Figure_{i}' for i in range(1,7)]+[f'Supplementary_{i}' for i in range(1,9)]
    missing_figures=[]
    for name in expected:
        unavailable=not (FIGURES/(name+'.png')).exists()
        if name in {'Figure_6','Supplementary_7'} and not native_completed():unavailable=True
        if unavailable:
            if not allow_pending or name not in {'Figure_6','Supplementary_7'}:raise FileNotFoundError(name)
            missing_figures.append(name)
    doc,template=fresh_template(template_path);caption=legends()
    groups=re.split(r'(?:\n\s*){2,}',source.strip())
    main_index=0;in_results=False;in_references=False;in_abstract=False
    def finish_section():
        nonlocal main_index
        if in_results and main_index:
            add_figure(doc,f'Figure_{main_index}',caption[f'Figure_{main_index}'],missing_figures)
    for text in groups:
        text=' '.join(text.splitlines()).strip()
        if text.startswith('### '):
            finish_section();main_index=0
            if in_results:
                titles=['Clone interpretation','In controlled GC models','*Plasmodium* reveals','Repeated human GC','Receptor-defined families','Captured-state readout']
                for i,title in enumerate(titles,1):
                    if text[4:].startswith(title):main_index=i
            p=doc.add_paragraph(style='Heading 3');inline(p,text[4:]);continue
        if text.startswith('## '):
            finish_section();main_index=0
            title=text[3:];in_results=title=='Results';in_references=title=='References';in_abstract=title=='Abstract'
            p=doc.add_paragraph(title,style='Heading 2');continue
        if text.startswith('# '):
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(18);p.paragraph_format.line_spacing=1.15
            run=p.add_run(text[2:]);run.bold=True;run.font.size=Pt(15);run.font.color.rgb=RGBColor.from_string('0F4761');continue
        p=doc.add_paragraph(style='EndNoteBibliography' if in_references else 'Normal')
        if not in_references:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        inline(p,text,italic=in_abstract)
    doc.add_page_break();doc.add_paragraph('Supplementary Information',style='Heading 1')
    for i in range(1,9):add_figure(doc,f'Supplementary_{i}',caption[f'Supplementary_{i}'],missing_figures)
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    audit=validate_docx(doc,template,output,order,missing_figures,template_path)
    output.with_suffix('.audit.json').write_text(json.dumps(audit,indent=2)+'\n')
    output.with_suffix('.ris').write_text(ris(refs,order))
    (PAPER/'MANUSCRIPT_draft_v2.md').write_text(source+'\n')
    if missing_figures:(PROJECT/'internal_validation/manuscript_preview_PENDING.md').write_text(source+'\n')
    print(json.dumps(audit,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--allow-pending',action='store_true')
    parser.add_argument('--output',type=Path)
    parser.add_argument('--template',type=Path,default=TEMPLATE)
    args=parser.parse_args()
    output=args.output or PROJECT/('Threadfin_MANUSCRIPT_GC_2026-10-06'+
        ('_REVIEW_PENDING_RBD' if args.allow_pending and not native_completed() else '')+'.docx')
    export(output,args.allow_pending,args.template)
